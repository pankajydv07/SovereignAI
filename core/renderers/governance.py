"""Official Deliverable Governance and Boundary Enforcement.

Prevents script-based document generation tools from producing official PSU
deliverables (approval notes, inspection summaries, calculations, cost sheets,
review decks) that must go through deterministic schema-validated renderers.
"""

import re
from pathlib import Path


class OfficialDeliverableViolationError(RuntimeError):
    """Raised when an attempt is made to generate an official deliverable via script execution."""

    def __init__(self, message: str) -> None:
        super().__init__(message)


OFFICIAL_FILENAME_PATTERNS = [
    "approvalnote",
    "inspectionsummary",
    "inspectionreport",
    "costsheet",
    "engineeringcalc",
    "engineeringcalculation",
    "calcsheet",
    "calculationderivation",
    "reviewdeck",
]

OFFICIAL_HEADER_PATTERNS = [
    r"approval\s+ladder\s*\(maker[\s_-]?checker\)",
    r"inter[\s_-]?office\s+memorandum\s*/\s*approval\s+note",
    r"technical\s+approval\s+note",
    r"inspection\s+sanction\s+affixed",
    r"deterministic\s+life\s+derivation\s*\(api\s*570\)",
]


def check_official_deliverable_boundary(
    output_filename: str, task_description: str, script_code: str | None = None
) -> None:
    """Enforce deterministic boundary prior to running script generation.

    Raises OfficialDeliverableViolationError if the requested document matches
    an official deliverable type.
    """
    ext = Path(output_filename).suffix.lower()
    if ext in (".docx", ".xlsx", ".pptx"):
        fn_clean = re.sub(r"[\W_]+", "", Path(output_filename).stem.lower())
        for pat in OFFICIAL_FILENAME_PATTERNS:
            if pat in fn_clean:
                raise OfficialDeliverableViolationError(
                    f"Attempted to produce official deliverable '{output_filename}' via script generation. "
                    "Official PSU deliverables (approval notes, inspection summaries, calculations, cost sheets, review decks) "
                    "are structurally forbidden in generate_document and must go through the deterministic 'render_deliverable' "
                    "tool with schema validation, approved organisation templates, and citation verification."
                )

    desc_lower = task_description.lower()
    if (
        ("approval note" in desc_lower or "inter-office memorandum" in desc_lower)
        and ("render" in desc_lower or "generate" in desc_lower or "create" in desc_lower)
    ):
        raise OfficialDeliverableViolationError(
            "Task description requests generating an official Approval Note via script. "
            "Official deliverables must be produced via the deterministic 'render_deliverable' tool."
        )



def scan_for_official_deliverable_structures(file_path: Path, output_format: str) -> None:
    """Scan pre-inspection/post-generation output for official PSU deliverable signatures."""
    if not file_path.exists():
        return

    extracted_text = ""
    fmt = output_format.lower().strip()

    if fmt == "docx":
        from docx import Document

        doc = Document(str(file_path))
        extracted_text = " ".join([p.text for p in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                extracted_text += " " + " ".join([c.text for c in row.cells])

    elif fmt == "md":
        extracted_text = file_path.read_text(encoding="utf-8")

    for pat in OFFICIAL_HEADER_PATTERNS:
        if re.search(pat, extracted_text, re.IGNORECASE):
            raise OfficialDeliverableViolationError(
                f"Generated document contains official PSU deliverable structure matching pattern '{pat}'. "
                "Official deliverables must be generated deterministically via 'render_deliverable'."
            )
