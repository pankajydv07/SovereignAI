"""Document Post-Processor & Provenance Injection.

Implements format-specific validation and provenance injection for:
- DOCX (python-docx header/footer paragraphs)
- XLSX (openpyxl _Attestation_Provenance worksheet + metadata)
- PPTX (python-pptx attestation closing slide)
- PDF (Per-page ReportLab transparent overlay merged via pypdf)
- Markdown (GFM frontmatter / provenance block)
"""

import re
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
from docx import Document
from docx.document import Document as DocumentClass
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches
from pptx.util import Pt as PptxPt
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas

from renderers.schemas import SystemProvenanceMetadata


class DocumentValidationError(Exception):
    """Raised when generated document fails structural or cardinality validation."""


class ProvenanceSpoofError(DocumentValidationError):
    """Raised when script attempted to author its own provenance/attestation block."""


def scan_for_provenance_spoofing(file_path: Path, output_format: str) -> None:
    """Scan pre-injection document content for illicit manual provenance/attestation markers."""
    suspicious_patterns = [
        r"run[\s_-]?id\s*:",
        r"\*\*\*\s*draft[\s_-]?warning\s*\*\*\*",
        r"provenance[\s_-]?attestation",
        r"models[\s_-]?used\s*:",
        r"min[\s_-]?conf(idence)?\s*:",
    ]
    extracted_text = ""

    fmt = output_format.lower().strip()
    if fmt == "docx":
        doc = Document(str(file_path))
        extracted_text = " ".join([p.text for p in doc.paragraphs])
        for section in doc.sections:
            extracted_text += " " + " ".join([p.text for p in section.footer.paragraphs])
            extracted_text += " " + " ".join([p.text for p in section.header.paragraphs])

    elif fmt == "xlsx":
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                extracted_text += " " + " ".join([str(v) for v in row if v is not None])

    elif fmt == "pptx":
        prs = Presentation(str(file_path))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    extracted_text += " " + shape.text_frame.text

    elif fmt == "pdf":
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                extracted_text += " " + txt

    elif fmt == "md":
        extracted_text = file_path.read_text(encoding="utf-8")

    for pat in suspicious_patterns:
        if re.search(pat, extracted_text, re.IGNORECASE):
            raise ProvenanceSpoofError(
                f"Manual provenance marker detected matching '{pat}'. "
                "The system injects official provenance automatically. "
                "Remove any manual footer/attestation blocks from the script."
            )


RAW_MARKDOWN_STRUCTURAL_PATTERNS = [
    (re.compile(r"^#{1,6}\s", re.MULTILINE), "Markdown heading marker"),
    (re.compile(r"\|[\s\-:]+\|"), "Markdown table delimiter row"),
    (re.compile(r"\*\*.+?\*\*"), "Markdown bold markers"),
]


def assert_zero_raw_markdown(text: str, doc_type: str = "document") -> None:
    """Ensure no raw structural Markdown was dumped directly into native document runs."""
    for pat, desc in RAW_MARKDOWN_STRUCTURAL_PATTERNS:
        match = pat.search(text)
        if match:
            raise DocumentValidationError(
                f"Generated {doc_type} contains unparsed {desc}: '{match.group(0)}'. "
                "The generator must emit native document elements (headings, tables, bold runs) "
                "rather than verbatim Markdown syntax."
            )


def validate_generated_document(
    file_path: Path,
    output_format: str,
    input_data: dict[str, Any] | list[Any] | None = None,
    expects_full_tabulation: bool = False,
) -> dict[str, Any]:
    """Validate structure, non-emptiness, and cardinality of the generated document."""
    if not file_path.exists():
        raise DocumentValidationError(f"Generated file does not exist: {file_path}")

    fmt = output_format.lower().strip()
    metrics: dict[str, Any] = {"format": fmt, "fileSizeBytes": file_path.stat().st_size}

    # First: Check for illicit provenance spoofing and official deliverable structures
    scan_for_provenance_spoofing(file_path, fmt)
    from renderers.governance import scan_for_official_deliverable_structures

    scan_for_official_deliverable_structures(file_path, fmt)

    expected_row_count = 0
    if isinstance(input_data, list):
        expected_row_count = len(input_data)
    elif isinstance(input_data, dict):
        for val in input_data.values():
            if isinstance(val, list):
                expected_row_count = max(expected_row_count, len(val))

    if fmt == "docx":
        try:
            doc = Document(str(file_path))
            p_count = len(doc.paragraphs)
            t_count = len(doc.tables)
            total_table_rows = sum(len(t.rows) for t in doc.tables)
            metrics.update(
                {
                    "paragraphs": p_count,
                    "tables": t_count,
                    "tableRows": total_table_rows,
                }
            )
            if p_count == 0 and t_count == 0:
                raise DocumentValidationError("DOCX has zero paragraphs and zero tables")

            full_text = "\n".join([p.text for p in doc.paragraphs])
            for t in doc.tables:
                for row in t.rows:
                    full_text += "\n" + " ".join([c.text for c in row.cells])
            assert_zero_raw_markdown(full_text, "DOCX")

            if expects_full_tabulation and expected_row_count > 0:
                if total_table_rows < expected_row_count:
                    raise DocumentValidationError(
                        f"DOCX table row count ({total_table_rows}) "
                        f"is less than expected input rows ({expected_row_count})"
                    )
        except Exception as exc:
            if isinstance(exc, DocumentValidationError):
                raise
            raise DocumentValidationError(f"Invalid DOCX document: {exc}") from exc

    elif fmt == "xlsx":
        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            sheet_names = wb.sheetnames
            metrics["sheets"] = sheet_names
            if not sheet_names:
                raise DocumentValidationError("XLSX contains no sheets")

            first_sheet = wb.active
            row_count = first_sheet.max_row if first_sheet else 0
            metrics["maxRows"] = row_count
            if row_count < 1:
                raise DocumentValidationError("XLSX active worksheet has no rows")

            sheet_text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    sheet_text += "\n" + " ".join([str(v) for v in row if v is not None])
            assert_zero_raw_markdown(sheet_text, "XLSX")

            if expects_full_tabulation and expected_row_count > 0:
                if row_count < expected_row_count:
                    raise DocumentValidationError(
                        f"XLSX row count ({row_count}) "
                        f"is less than expected input rows ({expected_row_count})"
                    )
        except Exception as exc:
            if isinstance(exc, DocumentValidationError):
                raise
            raise DocumentValidationError(f"Invalid XLSX workbook: {exc}") from exc

    elif fmt == "pptx":
        try:
            prs = Presentation(str(file_path))
            slide_count = len(prs.slides)
            metrics["slides"] = slide_count
            if slide_count < 1:
                raise DocumentValidationError("PPTX contains zero slides")

            prs_text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        prs_text += "\n" + shape.text_frame.text
            assert_zero_raw_markdown(prs_text, "PPTX")
        except Exception as exc:
            if isinstance(exc, DocumentValidationError):
                raise
            raise DocumentValidationError(f"Invalid PPTX presentation: {exc}") from exc

    elif fmt == "pdf":
        try:
            raw_bytes = file_path.read_bytes()
            if not raw_bytes.startswith(b"%PDF-"):
                raise DocumentValidationError("PDF missing standard %PDF- magic header")

            with pdfplumber.open(str(file_path)) as pdf:
                page_count = len(pdf.pages)
                metrics["pages"] = page_count
                if page_count < 1:
                    raise DocumentValidationError("PDF contains zero pages")

                total_chars = sum(len(p.extract_text() or "") for p in pdf.pages)
                metrics["charCount"] = total_chars
                if total_chars < 50:
                    raise DocumentValidationError("PDF has insufficient extracted text content")
        except Exception as exc:
            if isinstance(exc, DocumentValidationError):
                raise
            raise DocumentValidationError(f"Invalid PDF document: {exc}") from exc

    elif fmt == "md":
        content = file_path.read_text(encoding="utf-8")
        metrics["charCount"] = len(content)
        if len(content.strip()) < 10:
            raise DocumentValidationError("Markdown document is empty or too short")

    else:
        raise DocumentValidationError(f"Unsupported document format: {fmt}")

    return metrics


def inject_system_provenance(
    file_path: Path,
    output_format: str,
    prov: SystemProvenanceMetadata,
) -> Path:
    """Inject authoritative system provenance attestation into the validated document."""
    from renderers.provenance_injector import (
        inject_docx_provenance,
        inject_md_provenance,
        inject_pdf_provenance,
        inject_pptx_provenance,
        inject_xlsx_provenance,
    )

    fmt = output_format.lower().strip()
    if fmt == "docx":
        inject_docx_provenance(file_path, prov)
    elif fmt == "xlsx":
        inject_xlsx_provenance(file_path, prov)
    elif fmt == "pptx":
        inject_pptx_provenance(file_path, prov)
    elif fmt == "pdf":
        inject_pdf_provenance(file_path, prov)
    elif fmt == "md":
        inject_md_provenance(file_path, prov)
    else:
        raise ValueError(f"Unsupported format for provenance injection: {fmt}")

    return file_path

