"""Unit tests for Multi-Format Rich Document Reader and fs_read integration."""

import tempfile
from pathlib import Path
import pytest
import openpyxl
import docx
import pptx

from tools.document_reader import (
    convert_document_to_markdown,
    read_csv_to_markdown,
    read_docx_to_markdown,
    read_excel_to_markdown,
    read_pptx_to_markdown,
)
from tools.fs_read import FsReadInput, FsReadTool
from tools.base import ToolContext


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as d:
        yield Path(d)


def test_excel_to_markdown_multi_sheet(temp_dir: Path):
    """Verify Excel workbook with multiple sheets converts to Markdown tables."""
    xlsx_path = temp_dir / "plant_data.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Pump Data
    ws1 = wb.active
    ws1.title = "Pumps"
    ws1.append(["Tag", "Type", "Flow_m3h", "Head_m"])
    ws1.append(["P-101A", "Centrifugal", 150.5, 45.0])
    ws1.append(["P-101B", "Centrifugal", 150.5, 45.0])

    # Sheet 2: Valves Data
    ws2 = wb.create_sheet(title="Valves")
    ws2.append(["Valve_Tag", "Rating_Class", "Size_inch"])
    ws2.append(["MOV-201", "300#", 6])
    ws2.append(["PRV-202", "600#", 4])

    wb.save(str(xlsx_path))
    wb.close()

    md = read_excel_to_markdown(xlsx_path)
    assert "# Workbook: plant_data.xlsx" in md
    assert "## Sheet: Pumps" in md
    assert "| Tag | Type | Flow_m3h | Head_m |" in md
    assert "| P-101A | Centrifugal | 150.5 |" in md
    assert "## Sheet: Valves" in md
    assert "| MOV-201 | 300# | 6 |" in md


def test_docx_to_markdown(temp_dir: Path):
    """Verify Word document with headings and tables converts to structured Markdown."""
    docx_path = temp_dir / "inspection.docx"
    doc = docx.Document()
    doc.add_heading("Refinery Inspection Report", level=1)
    doc.add_paragraph("Visual inspection completed for Crude Distillation Unit.")
    
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Component"
    t.rows[0].cells[1].text = "Condition"
    t.rows[1].cells[0].text = "Column C-101"
    t.rows[1].cells[1].text = "Satisfactory"

    doc.save(str(docx_path))

    md = read_docx_to_markdown(docx_path)
    assert "## Refinery Inspection Report" in md
    assert "Visual inspection completed for Crude Distillation Unit." in md
    assert "| Component | Condition |" in md
    assert "| Column C-101 | Satisfactory |" in md


def test_pptx_to_markdown(temp_dir: Path):
    """Verify PowerPoint presentation slides convert to Markdown."""
    pptx_path = temp_dir / "presentation.pptx"
    prs = pptx.Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    tx_box = slide.shapes.add_textbox(100, 100, 300, 200)
    tf = tx_box.text_frame
    tf.text = "Monthly Turnaround Review"
    p = tf.add_paragraph()
    p.text = "All turnaround activities completed on schedule."

    prs.save(str(pptx_path))

    md = read_pptx_to_markdown(pptx_path)
    assert "## Slide 1" in md
    assert "Monthly Turnaround Review" in md
    assert "All turnaround activities completed on schedule." in md


def test_csv_to_markdown(temp_dir: Path):
    """Verify CSV files convert to Markdown table."""
    csv_path = temp_dir / "records.csv"
    csv_path.write_text("ID,Name,Status\n101,Exchanger,Active\n102,Boiler,Maintenance\n", encoding="utf-8")

    md = read_csv_to_markdown(csv_path)
    assert "| ID | Name | Status |" in md
    assert "| 101 | Exchanger | Active |" in md


@pytest.mark.asyncio
async def test_fs_read_tool_on_excel(temp_dir: Path):
    """Verify fs_read tool transparently returns Markdown table and supports slicing for .xlsx."""
    xlsx_path = temp_dir / "inventory.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spares"
    ws.append(["Item_Code", "Description", "Quantity"])
    for i in range(1, 10):
        ws.append([f"SKU-{i:03d}", f"Gasket Type {i}", i * 5])
    wb.save(str(xlsx_path))
    wb.close()

    tool = FsReadTool()
    ctx = ToolContext(workspace_root=temp_dir)

    # Read whole file
    res = await tool.run(FsReadInput(path="inventory.xlsx"), ctx)
    assert res.success is True
    assert "| Item_Code | Description | Quantity |" in str(res.output)
    assert "| SKU-001 | Gasket Type 1 | 5 |" in str(res.output)

    # Read sliced range
    res_sliced = await tool.run(FsReadInput(path="inventory.xlsx", start_line=1, end_line=5), ctx)
    assert res_sliced.success is True
    assert len(str(res_sliced.output).splitlines()) <= 5
