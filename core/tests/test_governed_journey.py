"""Flagship Governed Journey Milestone Integration Test.

Proves the complete, air-gapped governed workflow end-to-end:
1. Scanned inspection report ingest & OCR/vision escalation
2. Structured typed findings & InspectionReportSummary
3. Low-confidence verification gate blocking unverified approval notes
4. Human checker bounding box crop verification
5. Deterministic PSU org template DOCX deliverable generation
6. Maker-checker separation of duties enforcement & approval
7. Tamper-evident cryptographic audit chain sealing
"""

import time
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from approval.service import ApprovalService, SeparationOfDutiesError
from audit.chain import AuditChainStore
from ingest.pipeline import DocumentIngestPipeline
from ingest.types import BoundingBox, ExtractedWord
from renderers.engine import DeliverableRenderEngine
from renderers.schemas import (
    ApprovalLadderRole,
    ApprovalNoteSchema,
    CitationRef,
    InspectionFinding,
    InspectionReportSummary,
    LowConfidenceFieldUnverifiedError,
    SubstantiveClaim,
    ThicknessReading,
)
from renderers.templates import TemplateManager
from storage.db import DatabaseManager
from tests.fixtures.synthetic_generator import generate_realistic_scanned_fixture


class MockVisionModelClient:
    """Mock vision model for handwriting and stamp region extraction."""

    async def extract_field(self, crop_img, region) -> dict[str, str]:
        return {
            "field_name": "inspector_handwritten_remark",
            "value": "Immediate weld overlay required during next turnaround.",
            "unit": None,
        }


@pytest.mark.asyncio
async def test_complete_governed_journey(tmp_path: Path) -> None:
    """End-to-end integration test proving the complete governed deliverable journey."""
    start_time = time.perf_counter()

    # Setup database and audit chain
    db_file = tmp_path / "governed_journey.db"
    db_mgr = DatabaseManager(db_file)
    async with db_mgr.connect() as conn:
        await db_mgr.initialize_schema(conn)
        audit_store = AuditChainStore(conn)

        # ---------------------------------------------------------------------
        # STAGE 1: Scanned Inspection Report Ingest & Vision Escalation
        # ---------------------------------------------------------------------
        scan_file = tmp_path / "scan_c101_inspection.png"
        gt = generate_realistic_scanned_fixture(scan_file)
        assert scan_file.exists()

        mock_words = [
            ExtractedWord(text="Equipment", bbox=BoundingBox(x0=0.1, y0=0.15, x1=0.2, y1=0.18), confidence=0.95, page=1),
            ExtractedWord(text="Tag:", bbox=BoundingBox(x0=0.21, y0=0.15, x1=0.25, y1=0.18), confidence=0.95, page=1),
            ExtractedWord(text="C-101", bbox=BoundingBox(x0=0.26, y0=0.15, x1=0.35, y1=0.18), confidence=0.98, page=1),
            ExtractedWord(text="Inspection", bbox=BoundingBox(x0=0.5, y0=0.15, x1=0.6, y1=0.18), confidence=0.95, page=1),
            ExtractedWord(text="Date:", bbox=BoundingBox(x0=0.61, y0=0.15, x1=0.65, y1=0.18), confidence=0.95, page=1),
            ExtractedWord(text="2026-09-07", bbox=BoundingBox(x0=0.66, y0=0.15, x1=0.78, y1=0.18), confidence=0.98, page=1),
            ExtractedWord(text="Location", bbox=BoundingBox(x0=0.1, y0=0.25, x1=0.2, y1=0.28), confidence=0.92, page=1),
            ExtractedWord(text="CML:", bbox=BoundingBox(x0=0.21, y0=0.25, x1=0.26, y1=0.28), confidence=0.92, page=1),
            ExtractedWord(text="CML-07", bbox=BoundingBox(x0=0.27, y0=0.25, x1=0.36, y1=0.28), confidence=0.96, page=1),
            ExtractedWord(text="Measured", bbox=BoundingBox(x0=0.4, y0=0.25, x1=0.5, y1=0.28), confidence=0.92, page=1),
            ExtractedWord(text="Thickness:", bbox=BoundingBox(x0=0.51, y0=0.25, x1=0.62, y1=0.28), confidence=0.92, page=1),
            ExtractedWord(text="8.2", bbox=BoundingBox(x0=0.63, y0=0.25, x1=0.68, y1=0.28), confidence=0.74, page=1),
            ExtractedWord(text="mm", bbox=BoundingBox(x0=0.69, y0=0.25, x1=0.73, y1=0.28), confidence=0.90, page=1),
            # Handwritten region with low confidence to trigger vision escalation
            ExtractedWord(text="Handwritten", bbox=BoundingBox(x0=0.1, y0=0.50, x1=0.3, y1=0.55), confidence=0.30, page=1),
            ExtractedWord(text="Remark", bbox=BoundingBox(x0=0.31, y0=0.50, x1=0.5, y1=0.55), confidence=0.25, page=1),
        ]

        mock_vision = MockVisionModelClient()
        from concurrent.futures import ThreadPoolExecutor
        from unittest.mock import patch
        
        executor = ThreadPoolExecutor(max_workers=1)
        pipeline = DocumentIngestPipeline(vision_client=mock_vision, executor=executor)

        with patch("ingest.pipeline.extract_words_tesseract", return_value=mock_words):
            try:
                ingest_res = await pipeline.ingest_document(scan_file)
            finally:
                pipeline.shutdown()
                executor.shutdown(wait=False)

        assert len(ingest_res.pages) >= 1
        page_fields = ingest_res.fields
        assert len(page_fields) > 0

        # Verify vision-escalated field has uncalibrated confidence (None) and requires verification
        vision_fields = [f for f in page_fields if f.extractor == "vision_llm"]
        assert len(vision_fields) >= 1
        v_field = vision_fields[0]
        assert v_field.confidence is None
        assert v_field.requires_verification is True
        assert v_field.bbox is not None

        await audit_store.append_record(
            run_id="run-journey-01",
            user_id="user_sharma",
            action_type="INGEST_EVENT",
            prompt="Ingest scanned inspection report for Crude Distillation Column C-101",
            plan=[{"step": 1, "action": "OCR & Vision Escalation"}],
            steps=[{"step": 1, "status": "COMPLETED", "fields_extracted": len(page_fields)}],
            documents_retrieved=[{"doc_id": "MRPL-INSP-2026-09", "title": "Inspection Report C-101"}],
            models_used=["glm-ocr", "qwen-vl:7b"],
            deliverables=[],
        )

        # ---------------------------------------------------------------------
        # STAGE 2: Structured Typed Inspection Findings
        # ---------------------------------------------------------------------
        finding = InspectionFinding(
            finding_id="FIND-C101-01",
            category="CORROSION",
            severity="HIGH",
            equipment_tag="C-101",
            description=gt["defect_description"],
            source_page=1,
            source_region_bbox=[0.35, 0.08, 0.42, 0.92],
            extracted_field_ids=["t_actual_cml07", "corrosion_rate_cml07"],
            citations=[
                CitationRef(
                    doc_id="MRPL-INSP-2026-09",
                    clause_or_section="p. 1, CML-07",
                    extracted_field_id="t_actual_cml07",
                    confidence=0.74,  # Unverified low confidence
                    is_verified=False,
                    source_page=1,
                )
            ],
            confidence=0.74,
            verification_status="UNVERIFIED",
            recommended_action=gt["handwritten_inspector_remark"],
            inspection_method="UT + Visual",
        )

        inspection_summary = InspectionReportSummary(
            asset_tag="C-101",
            equipment_tag="C-101",
            inspection_date="2026-09-07",
            inspection_method="Ultrasonic Testing (UT) + Visual",
            findings=[finding],
            thickness_readings=[
                ThicknessReading(
                    location="CML-07",
                    nominal_mm=12.0,
                    actual_mm=8.2,
                    min_required_mm=4.5,
                )
            ],
            ncrs=[gt["ncr_ref"]],
            recommendations=[
                SubstantiveClaim(
                    text="Execute weld overlay cladding and quarterly UT thickness monitoring.",
                    citations=[CitationRef(doc_id="API-570", clause_or_section="Section 7.1.2", confidence=1.0)],
                )
            ],
        )
        assert len(inspection_summary.findings) == 1
        assert inspection_summary.findings[0].category == "CORROSION"

        # ---------------------------------------------------------------------
        # STAGE 3: Low-Confidence Gate Blocks Unverified Approval Note
        # ---------------------------------------------------------------------
        draft_note = ApprovalNoteSchema(
            subject="Technical Approval Note: Column C-101 Inspection Sanction",
            reference=["MRPL/INSP/2026/09"],
            background="Ultrasonic survey detected localised wall thinning in shell course 3.",
            observations=[
                SubstantiveClaim(
                    text="Measured wall thickness of 8.2 mm at CML-07 indicates active internal thinning.",
                    citations=[
                        CitationRef(
                            doc_id="MRPL-INSP-2026-09",
                            extracted_field_id="t_actual_cml07",
                            confidence=0.74,
                            is_verified=False,  # Unverified!
                            source_page=1,
                        )
                    ],
                )
            ],
            financial_implication="INR 12,50,000 under Turnaround Maintenance budget head.",
            deviation=None,
            recommendation=[
                SubstantiveClaim(
                    text="Sanction 14.8 months extended operational run subject to quarterly UT monitoring.",
                    citations=[CitationRef(doc_id="API-570", clause_or_section="Section 7.1.2", confidence=1.0)],
                )
            ],
            approval_ladder=[
                ApprovalLadderRole(role="Prepared By", name="A. Sharma (Senior Inspection Engineer)", status="VERIFIED"),
                ApprovalLadderRole(role="Checked By", name="P. V. Kulkarni (Chief Manager - Mechanical)", status="PENDING"),
            ],
        )

        tm = TemplateManager(workspace_root=tmp_path)
        engine = DeliverableRenderEngine(template_manager=tm)
        out_docx = tmp_path / "approval_note_c101.docx"

        # Attempt to render unverified approval note -> MUST fail loudly
        with pytest.raises(LowConfidenceFieldUnverifiedError) as exc_info:
            engine.render(
                deliverable_type="approval_note",
                data=draft_note.model_dump(),
                run_id="run-journey-01",
                output_path=out_docx,
            )
        assert "t_actual_cml07" in str(exc_info.value)
        assert "Checker verification is required" in str(exc_info.value)

        # ---------------------------------------------------------------------
        # STAGE 4: Checker Provenance & Crop Verification
        # ---------------------------------------------------------------------
        # Human checker inspects the scan crop on the workbench and confirms provenance
        draft_note.observations[0].citations[0].is_verified = True

        await audit_store.append_record(
            run_id="run-journey-01",
            user_id="user_kulkarni",
            action_type="FIELD_VERIFICATION",
            prompt="Verify crop bounding box for CML-07 wall thickness (8.2 mm)",
            plan=[],
            steps=[{"field": "t_actual_cml07", "status": "VERIFIED", "verified_by": "user_kulkarni"}],
            documents_retrieved=[],
            models_used=[],
            deliverables=[],
        )

        # ---------------------------------------------------------------------
        # STAGE 5: Deterministic Org Template DOCX Rendering
        # ---------------------------------------------------------------------
        # Register mandatory approved organisation template
        tpl_dir = tmp_path / ".swaraj" / "templates"
        tpl_dir.mkdir(parents=True, exist_ok=True)
        tpl_file = tpl_dir / "approval_note.docx"
        doc = Document()
        doc.add_heading("{{ subject }}", level=1)
        doc.add_paragraph("Background: {{ background }}")
        doc.add_paragraph("Financial Implication: {{ financial_implication }}")
        doc.save(str(tpl_file))

        rendered_doc = engine.render(
            deliverable_type="approval_note",
            data=draft_note.model_dump(),
            run_id="run-journey-01",
            output_path=out_docx,
        )
        assert rendered_doc.exists()
        assert rendered_doc.stat().st_size > 0

        await audit_store.append_record(
            run_id="run-journey-01",
            user_id="user_sharma",
            action_type="DELIVERABLE_RENDERED",
            prompt="Render official approval note DOCX",
            plan=[],
            steps=[{"template": "approval_note.docx", "status": "SUCCESS"}],
            documents_retrieved=[],
            models_used=["qwen3-coder:30b"],
            deliverables=[{"path": str(rendered_doc), "type": "approval_note", "sha256": "mocksha256"}],
        )

        # ---------------------------------------------------------------------
        # STAGE 6: Maker-Checker Separation of Duties, Stamping & Approval
        # ---------------------------------------------------------------------
        mock_pdf = tmp_path / "approval_note_C101_2026_approved.pdf"
        mock_pdf.write_bytes(b"%PDF-1.4 Mock PSU Approval Note PDF")

        mock_conv = MagicMock()
        mock_conv.convert_docx_to_pdf.return_value = mock_pdf
        approval_service = ApprovalService(pdf_converter=mock_conv)

        # Maker cannot check own work -> MUST fail
        with pytest.raises(SeparationOfDutiesError) as sod_err:
            approval_service.approve_deliverable(
                deliverable_id="DELIV-C101-2026",
                maker_id="user_sharma",
                checker_id="user_sharma",  # Self-approval attempt!
                checker_name="A. Sharma",
                checker_designation="Senior Inspection Engineer",
                unverified_fields_count=0,
                uncited_claims_count=0,
                calc_verification_status="VERIFIED",
                reviewed_docx_path=rendered_doc,
                output_dir=tmp_path,
            )
        assert "cannot be the preparing user" in str(sod_err.value)

        # Independent checker approves
        app_result = approval_service.approve_deliverable(
            deliverable_id="DELIV-C101-2026",
            maker_id="user_sharma",
            checker_id="user_kulkarni",
            checker_name="P. V. Kulkarni",
            checker_designation="Chief Manager - Mechanical",
            unverified_fields_count=0,
            uncited_claims_count=0,
            calc_verification_status="VERIFIED",
            reviewed_docx_path=rendered_doc,
            output_dir=tmp_path,
        )
        assert app_result.status == "APPROVED"
        assert app_result.audit_record.checker_id == "user_kulkarni"
        assert app_result.reviewed_docx_hash is not None
        assert app_result.stamped_docx_hash is not None
        assert app_result.approved_pdf_hash is not None
        assert app_result.pdf_status == "COMPLETED"
        stamp_text = app_result.stamp_text

        # ---------------------------------------------------------------------
        # STAGE 7: Sealed Cryptographic Audit Record (Three-Artifact Provenance)
        # ---------------------------------------------------------------------
        await audit_store.append_record(
            run_id="run-journey-01",
            user_id="user_kulkarni",
            action_type="APPROVAL_EVENT",
            prompt="Sanction and sign approval note DELIV-C101-2026",
            plan=[],
            steps=[{"action": "APPROVED", "deliverable_id": "DELIV-C101-2026"}],
            documents_retrieved=[],
            models_used=[],
            deliverables=[
                {"path": str(rendered_doc), "type": "reviewed_docx", "sha256": app_result.reviewed_docx_hash},
                {"path": str(tmp_path / "approval_note_C101_2026_approved.docx"), "type": "stamped_docx", "sha256": app_result.stamped_docx_hash},
                {"path": str(mock_pdf), "type": "approved_pdf", "sha256": app_result.approved_pdf_hash},
            ],
            approval_event={
                "maker": "user_sharma",
                "checker": "user_kulkarni",
                "action": "APPROVED",
                "stamp": stamp_text,
                "reviewed_docx_hash": app_result.reviewed_docx_hash,
                "stamped_docx_hash": app_result.stamped_docx_hash,
                "approved_pdf_hash": app_result.approved_pdf_hash,
                "pdf_status": app_result.pdf_status,
            },
        )

        all_records = await audit_store.get_all_records()
        assert len(all_records) == 4

        # Verify cryptographic chain integrity
        verification = AuditChainStore.verify_chain(all_records)
        assert verification.is_valid is True
        assert verification.total_records == 4
        assert verification.failed_index is None

    duration_s = time.perf_counter() - start_time
    print(f"\n[BENCHMARK] Complete Governed Journey Duration: {duration_s:.3f}s (Target < 90s)")
    assert duration_s < 90.0
