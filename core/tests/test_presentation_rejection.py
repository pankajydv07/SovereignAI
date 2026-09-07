"""Test suite for presentation requests classification, generation, and governance."""

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
import pytest
from pptx import Presentation

from agent.budget import RunBudget, RunBudgetTracker
from agent.turn_loop import TurnLoop
from models.registry import ModelRegistry
from models.router import ModelRouter, TaskClass
from renderers.governance import (
    OfficialDeliverableViolationError,
    check_official_deliverable_boundary,
)
from renderers.schemas import CitationRef, DeckSlide, ReviewDeckSchema
from tools.base import ToolContext
from tools.generate_document import GenerateDocumentInput, GenerateDocumentTool
from tools.registry import ToolRegistry
from tools.render_deliverable import RenderDeliverableInput, RenderDeliverableTool


@pytest.mark.asyncio
async def test_presentation_requests_classify_as_official_drafting():
    """Verify presentation requests route to official_drafting without a separate task class."""
    registry = ModelRegistry()
    router = ModelRouter(registry)

    presentation_prompts = [
        "create a management review deck",
        "make slides from this report",
        "generate a PPTX presentation",
    ]

    for prompt in presentation_prompts:
        decision = await router.route(prompt)
        assert decision.task_class == TaskClass.OFFICIAL_DRAFTING, (
            f"Expected {TaskClass.OFFICIAL_DRAFTING} for prompt '{prompt}', got {decision.task_class}"
        )


@pytest.mark.asyncio
async def test_generate_document_pptx_success(tmp_path):
    """Verify generate_document generates a valid PPTX file from task description / markdown."""
    tool = GenerateDocumentTool()
    ctx = ToolContext(workspace_root=tmp_path)

    res = await tool.run(
        GenerateDocumentInput(
            taskDescription="# Turnaround Inspection Overview\n\n## Key Findings\n- Boiler B-201 tube corrosion\n- Column C-101 tray damage",
            outputFormat="pptx",
            outputFilename="turnaround_slides.pptx",
        ),
        ctx,
    )
    assert res.success
    assert Path(res.output.file_path).exists()
    prs = Presentation(res.output.file_path)
    assert len(prs.slides) >= 2
    assert "Turnaround Inspection Overview" in prs.slides[0].shapes.title.text


@pytest.mark.asyncio
async def test_render_deliverable_review_deck_success(tmp_path):
    """Verify render_deliverable successfully generates review_deck.pptx with citations and provenance."""
    tool = RenderDeliverableTool()
    ctx = ToolContext(workspace_root=tmp_path)

    schema_data = {
        "title": "Quarterly Plant Asset Review",
        "subtitle": "IOCL Panipat Refinery - Area 4",
        "slides": [
            {
                "slide_type": "metrics",
                "title": "Corrosion Monitoring Summary",
                "bullets": ["Critical thinning detected at Circuit 14", "Immediate UT scanning recommended"],
                "key_metrics": [{"CIRCUIT 14 MIN MM": "4.2 mm"}, {"RETIREMENT LIMIT": "3.8 mm"}],
                "citations": [{"doc_id": "INSP-2026-08", "clause_or_section": "Sec 4.2"}],
            }
        ],
    }

    res = await tool.run(
        RenderDeliverableInput(
            deliverableType="review_deck",
            data=schema_data,
            outputFilename="asset_review.pptx",
        ),
        ctx,
    )
    assert res.success
    assert Path(res.output.file_path).exists()
    prs = Presentation(res.output.file_path)
    assert len(prs.slides) >= 3  # Title, Content, and Final Provenance Slide
    assert "Quarterly Plant Asset Review" in prs.slides[0].shapes.title.text


def test_governance_boundary_protects_official_review_deck():
    """Verify governance boundary check enforces render_deliverable for official review deck names."""
    with pytest.raises(OfficialDeliverableViolationError, match=r"Official PSU deliverables"):
        check_official_deliverable_boundary(
            output_filename="management_reviewdeck.pptx",
            task_description="Render slides for management review",
        )


@pytest.mark.asyncio
async def test_turn_loop_multi_turn_format_isolation(tmp_path):
    """Verify XLSX request in a session with past PPTX turn isolates format to XLSX."""
    registry = ToolRegistry()
    registry.register(GenerateDocumentTool())

    mock_ollama = MagicMock()
    async def mock_stream_chat(*args, **kwargs):
        yield {"message": {"content": "Here is the spreadsheet:\n| Component | Thickness |\n| Shell | 14.2 mm |"}, "done": True}
    mock_ollama.stream_chat = mock_stream_chat

    turn_loop = TurnLoop(
        ollama_client=mock_ollama,
        tool_registry=registry,
        budget_tracker=RunBudgetTracker(budget=RunBudget(max_steps=5)),
        permission_requester=AsyncMock(return_value=("allow_once", None)),
    )

    history = [
        {"role": "user", "content": "create a presentation on refinery inspection"},
        {"role": "assistant", "content": "Generated `presentation.pptx` in the workspace with system provenance attestation."},
        {"role": "user", "content": "now generate an xlsx sheet of pipe thickness audit"},
    ]

    reason, updated_messages = await turn_loop.run_step(
        session_id="test-session",
        model_tag="test-model",
        messages=history,
        task_class="official_drafting",
        tool_context=ToolContext(workspace_root=tmp_path),
    )

    assert (tmp_path / "data.xlsx").exists()
    assert any("data.xlsx" in m.get("content", "") for m in updated_messages if m.get("role") == "assistant")


