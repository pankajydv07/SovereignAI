"""Script to generate PSU Word templates for SWARAJ with strict design lock.

Generates:
- core/templates/approval_note.docx
- core/templates/inspection_summary.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _setup_section(s):
    s.page_width = Mm(210)
    s.page_height = Mm(297)
    s.top_margin = Mm(25)
    s.bottom_margin = Mm(25)
    s.left_margin = Mm(30)
    s.right_margin = Mm(20)
    s.header_distance = Mm(12)
    s.footer_distance = Mm(12)

    # Header
    hdr = s.header
    p_hdr = hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_hdr.add_run("CONFIDENTIAL — FOR INTERNAL USE ONLY\n")
    r_c.font.name = "IBM Plex Mono"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)
    r_d = p_hdr.add_run("DRAFT — REQUIRES APPROVAL BY COMPETENT AUTHORITY")
    r_d.bold = True
    r_d.font.name = "IBM Plex Mono"
    r_d.font.size = Pt(9.0)
    r_d.font.color.rgb = RGBColor(0xF5, 0x9E, 0x0B)

    # Footer
    ftr = s.footer
    p_ftr = ftr.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_f = p_ftr.add_run("Run ID: {{ provenance.run_id }} | Models: {{ provenance.models_used | join(', ') }}")
    r_f.font.name = "IBM Plex Mono"
    r_f.font.size = Pt(7.5)
    r_f.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)


def _apply_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)


def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    else:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    return h


def create_approval_note_template(out_dir: Path) -> Path:
    doc = Document()
    _setup_section(doc.sections[0])
    _apply_styles(doc)

    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_top.add_run("INDIAN OIL CORPORATION LIMITED\nREFINERIES DIVISION — PANIPAT REFINERY\nNOTE FOR APPROVAL")
    r_t.bold = True
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(13)
    r_t.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    _add_heading(doc, "1. SUBJECT", 1)
    doc.add_paragraph("{{ subject }}")

    _add_heading(doc, "2. REFERENCE & INITIATION", 1)
    doc.add_paragraph("{% for ref in reference %}{{ ref }}{% if not loop.last %} | {% endif %}{% endfor %}")

    _add_heading(doc, "3. BACKGROUND & OPERATIONAL CONTEXT", 1)
    doc.add_paragraph("{{ background }}")

    _add_heading(doc, "4. TECHNICAL OBSERVATIONS & DEFECTS", 1)
    doc.add_paragraph("{% for obs in observations %}• {{ obs.text if obs.text is defined else obs }}\n{% endfor %}")

    _add_heading(doc, "5. FINANCIAL & STATUTORY IMPLICATIONS", 1)
    doc.add_paragraph("{{ financial_implication }}")

    _add_heading(doc, "6. RECOMMENDATIONS FOR SANCTION", 1)
    doc.add_paragraph("{% for rec in recommendation %}• {{ rec.text if rec.text is defined else rec }}\n{% endfor %}")

    _add_heading(doc, "7. COMPETENT AUTHORITY APPROVAL LADDER", 1)
    tbl = doc.add_table(rows=4, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    hdr[0].text = "Role / Stage"
    hdr[1].text = "Officer Name & Designation"
    hdr[2].text = "Action Status"

    for c in hdr:
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "1F2328")
        tcPr.append(shd)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tbl.rows[1].cells[0].text = "{%tr for item in approval_ladder %}"
    
    row2 = tbl.rows[2].cells
    row2[0].text = "{{ item.role }}"
    row2[1].text = "{{ item.name }}"
    row2[2].text = "{{ item.status }}"

    tbl.rows[3].cells[0].text = "{%tr endfor %}"

    out_file = out_dir / "approval_note.docx"
    doc.save(str(out_file))
    return out_file


def create_inspection_summary_template(out_dir: Path) -> Path:
    doc = Document()
    _setup_section(doc.sections[0])
    _apply_styles(doc)

    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_top.add_run("INDIAN OIL CORPORATION LIMITED\nASSET INTEGRITY & NON-DESTRUCTIVE TESTING REPORT\nSUMMARY OF INSPECTION FINDINGS")
    r_t.bold = True
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(13)
    r_t.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    _add_heading(doc, "1. ASSET IDENTIFICATION & METADATA", 1)
    p_m = doc.add_paragraph()
    p_m.add_run("Asset Tag: {{ asset_tag }} | Date: {{ inspection_date }} | Method: {{ inspection_method }}").bold = True

    _add_heading(doc, "2. THICKNESS SURVEY READINGS", 1)
    tbl = doc.add_table(rows=4, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    hdr[0].text = "Location / Circuit"
    hdr[1].text = "Nominal (mm)"
    hdr[2].text = "Actual (mm)"
    hdr[3].text = "Min Required (mm)"

    for c in hdr:
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "1F2328")
        tcPr.append(shd)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tbl.rows[1].cells[0].text = "{%tr for t in thickness_readings %}"

    row2 = tbl.rows[2].cells
    row2[0].text = "{{ t.location }}"
    row2[1].text = "{{ '%.2f'|format(t.nominal_mm) }}"
    row2[2].text = "{{ '%.2f'|format(t.actual_mm) }}"
    row2[3].text = "{{ '%.2f'|format(t.min_required_mm) }}"

    tbl.rows[3].cells[0].text = "{%tr endfor %}"

    _add_heading(doc, "3. OBSERVED DEFECTS & FINDINGS", 1)
    doc.add_paragraph("{% for d in observed_defects %}• [{{ d.severity }}] {{ d.description }}\n{% endfor %}")

    _add_heading(doc, "4. ENGINEERING RECOMMENDATIONS", 1)
    doc.add_paragraph("{% for r in recommendations %}• {{ r.text if r.text is defined else r }}\n{% endfor %}")

    out_file = out_dir / "inspection_summary.docx"
    doc.save(str(out_file))
    return out_file


if __name__ == "__main__":
    templates_dir = Path("core/templates")
    templates_dir.mkdir(parents=True, exist_ok=True)
    create_approval_note_template(templates_dir)
    create_inspection_summary_template(templates_dir)
    print("Templates generated successfully in core/templates/")
